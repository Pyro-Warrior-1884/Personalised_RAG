"""Shared helpers for the test suite."""

from __future__ import annotations

import hashlib
import math
import re
from pathlib import Path


class FakeEmbeddings:
    """A deterministic, dependency free embedding model.

    The vector of a text is built from hashed tokens, so two texts that share
    words end up with a positive cosine similarity.  This keeps the tests
    hermetic: no model is downloaded and no network call is made, while the
    retrieval behaviour that matters (the right chunk is returned) is still
    exercised end to end.
    """

    dimension = 96

    def _embed(self, text: str) -> list[float]:
        vector = [0.0] * self.dimension
        for token in re.findall(r"[a-z0-9]+", (text or "").lower()):
            digest = hashlib.md5(token.encode("utf-8")).digest()
            for offset in range(0, 16, 2):
                index = ((digest[offset] << 8) | digest[offset + 1]) % self.dimension
                vector[index] += 1.0
        norm = math.sqrt(sum(value * value for value in vector)) or 1.0
        return [value / norm for value in vector]

    def embed_documents(self, texts: list[str]) -> list[list[float]]:
        return [self._embed(text) for text in texts]

    def embed_query(self, text: str) -> list[float]:
        return self._embed(text)

    async def aembed_documents(self, texts: list[str]) -> list[list[float]]:
        return self.embed_documents(texts)

    async def aembed_query(self, text: str) -> list[float]:
        return self.embed_query(text)


def make_pdf(path: Path, pages: list[str]) -> Path:
    """Write a small PDF containing the given page texts."""
    import pymupdf

    document = pymupdf.open()
    for text in pages:
        page = document.new_page()
        page.insert_text((72, 100), text[:120], fontsize=11)
    document.save(str(path))
    document.close()
    return path


def make_docx(path: Path, paragraphs: list[str]) -> Path:
    """Write a small Word document."""
    import docx

    document = docx.Document()
    for paragraph in paragraphs:
        document.add_paragraph(paragraph)
    document.save(str(path))
    return path


def make_xlsx(path: Path, rows: list[list[str]]) -> Path:
    """Write a small Excel workbook with a single sheet."""
    from openpyxl import Workbook

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Data"
    for row in rows:
        sheet.append(row)
    workbook.save(str(path))
    return path


def make_html(path: Path, title: str, body: str) -> Path:
    """Write a small HTML file."""
    path.write_text(
        f"<!doctype html><html><head><title>{title}</title>"
        f"<style>.x{{color:red}}</style></head>"
        f"<body><nav>Menu</nav><h1>{title}</h1><p>{body}</p>"
        f"<script>alert('x')</script></body></html>",
        encoding="utf-8",
    )
    return path
