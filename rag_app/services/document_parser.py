"""Format specific document parsing.

Every parser returns LangChain ``Document`` objects that carry metadata used
later for citations (relative path, page number, sheet name, row range ...).
Parsing never raises for a single bad file: a structured error is returned
so the ingestion job can record the failure and carry on.
"""

from __future__ import annotations

import csv
import io
import logging
import re
from dataclasses import dataclass, field
from pathlib import Path

from langchain_core.documents import Document

from .file_discovery import file_type_for, is_supported_file

logger = logging.getLogger(__name__)

MAX_EXCEL_ROWS = 5000
MAX_CSV_ROWS = 5000
MAX_TEXT_CHARS = 2_000_000
HTML_NOISE_TAGS = ("script", "style", "noscript", "nav", "header", "footer", "aside", "form", "svg")


@dataclass
class ParseResult:
    """The outcome of parsing a single file."""

    documents: list[Document] = field(default_factory=list)
    error: str | None = None
    reason: str | None = None
    file_type: str = ""

    @property
    def ok(self) -> bool:
        return self.error is None and bool(self.documents)


class DocumentParseError(Exception):
    """Raised internally when a file cannot be parsed."""

    def __init__(self, message: str, reason: str = "corrupt") -> None:
        super().__init__(message)
        self.reason = reason


def _collapse_blank_lines(text: str) -> str:
    """Collapse runs of blank lines and strip trailing whitespace."""
    lines = [line.rstrip() for line in text.replace("\r\n", "\n").replace("\r", "\n").split("\n")]
    cleaned: list[str] = []
    blank = 0
    for line in lines:
        if not line.strip():
            blank += 1
            if blank > 1:
                continue
        else:
            blank = 0
        cleaned.append(line)
    return "\n".join(cleaned).strip()


def _truncate(text: str, limit: int = MAX_TEXT_CHARS) -> str:
    if len(text) <= limit:
        return text
    return text[:limit] + "\n\n[... content truncated ...]"


# ---------------------------------------------------------------------------
# PDF
# ---------------------------------------------------------------------------
def load_pdf(path: Path, relative_path: str) -> list[Document]:
    """Extract page level text from a PDF using PyMuPDF."""
    try:
        import pymupdf  # type: ignore
    except ImportError:  # pragma: no cover - older wheels expose ``fitz``
        try:
            import fitz as pymupdf  # type: ignore
        except ImportError as exc:  # pragma: no cover
            raise DocumentParseError(
                "PyMuPDF is not installed, so PDF files cannot be parsed.", "unsupported"
            ) from exc

    try:
        document = pymupdf.open(str(path))
    except Exception as exc:  # noqa: BLE001 - any failure means unreadable PDF
        message = str(exc).lower()
        if "password" in message or "encrypt" in message or "authentication" in message:
            raise DocumentParseError("The PDF is encrypted or password protected.", "encrypted") from exc
        raise DocumentParseError("The PDF could not be opened (it may be corrupt).", "corrupt") from exc

    documents: list[Document] = []
    try:
        if getattr(document, "needs_pass", False):
            raise DocumentParseError("The PDF is encrypted or password protected.", "encrypted")
        for page_index in range(document.page_count):
            try:
                page = document.load_page(page_index)
                text = page.get_text("text") or ""
            except Exception as exc:  # noqa: BLE001
                logger.warning("Could not read page %s of %s: %s", page_index + 1, relative_path, exc)
                continue
            text = _collapse_blank_lines(text)
            if not text:
                continue
            documents.append(
                Document(
                    page_content=text,
                    metadata={
                        "relative_path": relative_path,
                        "file_type": "pdf",
                        "page_number": page_index + 1,
                        "source": relative_path,
                    },
                )
            )
    finally:
        try:
            document.close()
        except Exception:  # pragma: no cover - defensive
            pass

    if not documents:
        raise DocumentParseError(
            "No extractable text was found (the PDF may be a scanned image).", "empty"
        )
    return documents


# ---------------------------------------------------------------------------
# Excel
# ---------------------------------------------------------------------------
def _format_cell(value) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value)


def load_excel(path: Path, relative_path: str) -> list[Document]:
    """Extract workbook content keeping sheet and row context."""
    try:
        from openpyxl import load_workbook
    except ImportError as exc:  # pragma: no cover
        raise DocumentParseError("openpyxl is not installed.", "unsupported") from exc

    try:
        workbook = load_workbook(filename=str(path), read_only=True, data_only=True)
    except Exception as exc:  # noqa: BLE001
        raise DocumentParseError(
            "The workbook could not be opened (it may be corrupt or password protected).",
            "corrupt",
        ) from exc

    documents: list[Document] = []
    try:
        for sheet_name in workbook.sheetnames:
            sheet = workbook[sheet_name]
            rows: list[str] = []
            header: list[str] = []
            row_count = 0
            for index, row in enumerate(sheet.iter_rows(values_only=True), start=1):
                if index == 1:
                    header = [_format_cell(cell) for cell in row]
                cells = [_format_cell(cell) for cell in row]
                if not any(cell.strip() for cell in cells):
                    continue
                row_count = index
                if index == 1:
                    rows.append(" | ".join(cells))
                else:
                    rows.append(f"Row {index}: " + " | ".join(cells))
                if index > MAX_EXCEL_ROWS:
                    break
            if not rows:
                continue
            header_line = ""
            if header and any(item.strip() for item in header):
                header_line = "Columns: " + " | ".join(header) + "\n"
            text = _collapse_blank_lines(f"{header_line}\n" + "\n".join(rows))
            if not text:
                continue
            documents.append(
                Document(
                    page_content=_truncate(text),
                    metadata={
                        "relative_path": relative_path,
                        "file_type": "excel",
                        "workbook": Path(relative_path).name,
                        "sheet": sheet_name,
                        "row_start": 1,
                        "row_end": row_count,
                        "source": relative_path,
                    },
                )
            )
    finally:
        try:
            workbook.close()
        except Exception:  # pragma: no cover - defensive
            pass

    if not documents:
        raise DocumentParseError("The workbook contains no readable cells.", "empty")
    return documents


# ---------------------------------------------------------------------------
# HTML
# ---------------------------------------------------------------------------
def load_html(path: Path, relative_path: str) -> list[Document]:
    """Extract readable text from an HTML file using BeautifulSoup."""
    try:
        from bs4 import BeautifulSoup
    except ImportError as exc:  # pragma: no cover
        raise DocumentParseError("BeautifulSoup is not installed.", "unsupported") from exc

    try:
        raw = path.read_bytes()
    except OSError as exc:
        raise DocumentParseError("The file could not be read.", "corrupt") from exc

    html = raw.decode("utf-8", errors="replace")
    try:
        soup = BeautifulSoup(html, "html.parser")
    except Exception as exc:  # noqa: BLE001
        raise DocumentParseError("The HTML could not be parsed.", "corrupt") from exc

    for tag in soup(list(HTML_NOISE_TAGS)):
        tag.decompose()

    title = ""
    if soup.title and soup.title.string:
        title = str(soup.title.string).strip()
    elif soup.h1:
        title = soup.h1.get_text(strip=True)

    text = soup.get_text(separator="\n")
    text = _collapse_blank_lines(text)
    if not text:
        raise DocumentParseError("The HTML file contains no readable text.", "empty")

    metadata = {
        "relative_path": relative_path,
        "file_type": "html",
        "source": relative_path,
    }
    if title:
        metadata["page_title"] = title[:300]
    return [Document(page_content=_truncate(text), metadata=metadata)]


# ---------------------------------------------------------------------------
# DOCX
# ---------------------------------------------------------------------------
def load_docx(path: Path, relative_path: str) -> list[Document]:
    """Extract paragraphs and tables from a Word document."""
    try:
        import docx  # type: ignore
    except ImportError as exc:  # pragma: no cover
        raise DocumentParseError("python-docx is not installed.", "unsupported") from exc

    try:
        document = docx.Document(str(path))
    except Exception as exc:  # noqa: BLE001
        message = str(exc).lower()
        if "password" in message or "encrypt" in message:
            raise DocumentParseError(
                "The Word document is encrypted or password protected.", "encrypted"
            ) from exc
        raise DocumentParseError("The Word document could not be opened.", "corrupt") from exc

    parts: list[str] = []
    for index, paragraph in enumerate(document.paragraphs, start=1):
        text = (paragraph.text or "").strip()
        if text:
            parts.append(f"Paragraph {index}: {text}")

    for table_index, table in enumerate(document.tables, start=1):
        rows = []
        for row in table.rows:
            cells = [(cell.text or "").strip() for cell in row.cells]
            if any(cells):
                rows.append(" | ".join(cells))
        if rows:
            parts.append(f"Table {table_index}:\n" + "\n".join(rows))

    text = _collapse_blank_lines("\n".join(parts))
    if not text:
        raise DocumentParseError("The Word document contains no readable text.", "empty")

    return [
        Document(
            page_content=_truncate(text),
            metadata={
                "relative_path": relative_path,
                "file_type": "word",
                "source": relative_path,
            },
        )
    ]


# ---------------------------------------------------------------------------
# CSV
# ---------------------------------------------------------------------------
def load_csv(path: Path, relative_path: str) -> list[Document]:
    """Read a CSV file preserving the header and row numbers."""
    try:
        raw = path.read_bytes()
    except OSError as exc:
        raise DocumentParseError("The file could not be read.", "corrupt") from exc

    text = raw.decode("utf-8-sig", errors="replace")
    try:
        dialect = csv.Sniffer().sniff(text[:4096], delimiters=",;|\t")
    except csv.Error:
        dialect = csv.excel

    reader = csv.reader(io.StringIO(text), dialect)
    lines: list[str] = []
    header: list[str] = []
    row_count = 0
    try:
        for index, row in enumerate(reader, start=1):
            if index == 1:
                header = [cell.strip() for cell in row]
                lines.append(" | ".join(header))
                continue
            if not any(cell.strip() for cell in row):
                continue
            row_count = index
            lines.append(f"Row {index}: " + " | ".join(cell.strip() for cell in row))
            if index > MAX_CSV_ROWS:
                break
    except csv.Error as exc:
        raise DocumentParseError("The CSV could not be parsed.", "corrupt") from exc

    body = _collapse_blank_lines("\n".join(lines))
    if not body:
        raise DocumentParseError("The CSV contains no rows.", "empty")

    metadata = {
        "relative_path": relative_path,
        "file_type": "csv",
        "row_start": 1,
        "row_end": row_count,
        "source": relative_path,
    }
    if header and any(header):
        metadata["columns"] = " | ".join(header[:50])[:1000]
    return [Document(page_content=_truncate(body), metadata=metadata)]


# ---------------------------------------------------------------------------
# Plain text and source code
# ---------------------------------------------------------------------------
def _read_text(path: Path) -> str:
    try:
        raw = path.read_bytes()
    except OSError as exc:
        raise DocumentParseError("The file could not be read.", "corrupt") from exc
    for encoding in ("utf-8-sig", "utf-8", "cp1252", "latin-1"):
        try:
            return raw.decode(encoding)
        except UnicodeDecodeError:
            continue
    return raw.decode("utf-8", errors="replace")


def load_text_or_code(path: Path, relative_path: str) -> list[Document]:
    """Read a plain text, Markdown or source code file."""
    text = _read_text(path)
    if not text.strip():
        raise DocumentParseError("The file is empty.", "empty")

    file_type = file_type_for(path)
    metadata = {
        "relative_path": relative_path,
        "file_type": file_type,
        "source": relative_path,
    }
    if file_type == "code":
        metadata["language"] = path.suffix.lower().lstrip(".")

    if file_type == "markdown":
        headings = re.findall(r"^#{1,6}\s+(.+)$", text, flags=re.MULTILINE)[:20]
        if headings:
            metadata["headings"] = " | ".join(headings)[:1000]

    return [Document(page_content=_truncate(text), metadata=metadata)]


# ---------------------------------------------------------------------------
# Entry point
# ---------------------------------------------------------------------------
def load_file(path: Path, relative_path: str | None = None) -> ParseResult:
    """Parse a single file into LangChain documents.

    Never raises: problems are reported through the returned ``ParseResult``.
    """
    path = Path(path)
    relative_path = relative_path or path.name
    file_type = file_type_for(path)

    if not path.exists():
        return ParseResult(error=f"File not found: {relative_path}", reason="missing", file_type=file_type)
    if not path.is_file():
        return ParseResult(error=f"Not a file: {relative_path}", reason="missing", file_type=file_type)
    if not is_supported_file(path):
        return ParseResult(
            error=f"Unsupported file type: {path.suffix or 'unknown'}", reason="unsupported", file_type=file_type
        )
    if path.stat().st_size == 0:
        return ParseResult(error=f"The file is empty: {relative_path}", reason="empty", file_type=file_type)

    try:
        suffix = path.suffix.lower()
        if suffix == ".pdf":
            documents = load_pdf(path, relative_path)
        elif suffix in {".xlsx", ".xlsm"}:
            documents = load_excel(path, relative_path)
        elif suffix in {".html", ".htm"}:
            documents = load_html(path, relative_path)
        elif suffix == ".docx":
            documents = load_docx(path, relative_path)
        elif suffix == ".csv":
            documents = load_csv(path, relative_path)
        else:
            documents = load_text_or_code(path, relative_path)
    except DocumentParseError as exc:
        return ParseResult(error=str(exc), reason=exc.reason, file_type=file_type)
    except Exception as exc:  # noqa: BLE001 - never let one file break the batch
        logger.warning("Unexpected parse failure for %s: %s", relative_path, exc)
        return ParseResult(
            error=f"The file could not be parsed: {type(exc).__name__}", reason="corrupt", file_type=file_type
        )

    if not documents:
        return ParseResult(
            error=f"No readable content was found in {relative_path}", reason="empty", file_type=file_type
        )
    return ParseResult(documents=documents, file_type=file_type)
